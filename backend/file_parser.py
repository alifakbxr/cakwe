import io
from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Mengekstrak seluruh teks dari file PDF secara in-memory."""
    text_content = []
    reader = PdfReader(io.BytesIO(file_bytes))
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            text_content.append(extracted)
    return "\n".join(text_content).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Mengekstrak seluruh teks dari file DOCX secara in-memory."""
    doc = Document(io.BytesIO(file_bytes))
    text_content = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_content.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_content.append(" | ".join(row_text))

    return "\n".join(text_content).strip()


def parse_resume_document(filename: str, file_bytes: bytes) -> str:
    """Deteksi tipe file dan ekstrak teksnya."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        return extract_text_from_docx(file_bytes)
    elif lower_name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Format file tidak didukung. Harap unggah file PDF, DOCX, atau TXT.")
