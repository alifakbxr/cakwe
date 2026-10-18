import os
import io
import json
from google import genai
from google.genai import types
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from dotenv import load_dotenv

from schemas import (
    FullResumeData,
    OptimizedResumeResult,
    JobRequisition,
    ManualProfileBuildRequest,
)

load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")


def build_expert_ats_prompt() -> str:
    return """Anda adalah Principal Technical Recruiter dan Pakar Resume Engineering berstandar Harvard & MIT Career Services.
Tujuan Anda: Mengubah profil/resume kandidat menjadi resume ATS kelas dunia yang sangat relevan, padat, profesional, dan meyakinkan untuk lowongan target.

PEDOMAN KONTEN & FORMAT STANDAR HARVARD/MIT:
1. STRUKTUR SATU KOLOM (SINGLE-COLUMN ATS STANDARD):
   - Dokumen harus bersih, mudah dibaca mesin ATS, tanpa tabel kompleks, ikon gambar, atau grafik warna-warni.
   - Urutan seksi: Kontak -> Professional Summary -> Technical Skills -> Professional Experience -> Education -> Certifications (bila ada).

2. PROFESSIONAL SUMMARY BERBOBOT:
   - Tulis ringkasan 2–3 kalimat padat yang menyoroti: identitas profesional kandidat, keahlian utama yang cocok dengan lowongan target, dan nilai tambah yang siap diberikan kepada tim/perusahaan.
   - Hindari kata-kata klise kosong seperti "hard worker", "passionate team player". Gunakan pencapaian nyata dan lingkup keahlian.

3. PENGALAMAN KERJA DENGAN FORMULA DAMPAK GOOGLE (X-Y-Z):
   - Format: "Accomplished [X] as measured by [Y], by doing [Z]"
   - Awali setiap bullet point dengan Action Verb lampau yang kuat (mis: Architected, Spearheaded, Refactored, Accelerated, Deployed, Engineered, Streamlined).
   - Tunjukkan konteks teknis: sebutkan tools, bahasa, framework, atau arsitektur yang dipakai saat menyelesaikan masalah tersebut.
   - FOKUS HASIL/DAMPAK: Jelaskan peningkatan performa, efisiensi waktu, reduksi latensi, stabilitas sistem, atau kepuasan pengguna.

4. INTEGRITAS FAKTA & ANTI-HALUSINASI (ATURAN MUTLAK):
   - JANGAN PERNAH mengarang perusahaan, jabatan, durasi kerja, gelar, atau sertifikasi yang tidak ada di profil kandidat.
   - Jika metrik angka tidak dinyatakan oleh kandidat, gunakan placeholder terarah seperti [X%] atau [N ms] agar kandidat dapat mengisinya dengan angka faktual mereka, ATAU fokus pada pencapaian kualitatif nyata (mis: zero-downtime deployment, test coverage improvement).

5. COVER LETTER SPESIFIK & PERSUASIF:
   - Buat cover letter 3–4 paragraf yang profesional dan terarah:
     * Paragraf 1: Posisi yang dituju, ketertarikan pada misi perusahaan, dan ringkasan nilai jual utama.
     * Paragraf 2-3: Bukti konkrit 1-2 pencapaian paling relevan dari latar belakang kandidat yang langsung menjawab kebutuhan spesifik lowongan.
     * Paragraf 4: Penutup profesional, kesiapan wawancara, dan apresiasi.
"""


async def optimize_resume_mode_a(raw_resume_text: str, job: JobRequisition) -> OptimizedResumeResult:
    """Mode A: Optimasi resume yang diunggah / dipaste."""
    if not GEMINI_API_KEY:
        raise ValueError("GEMINI_API_KEY belum disetel di .env")

    client = genai.Client(api_key=GEMINI_API_KEY)

    prompt = f"""LOWONGAN PEKERJAAN TARGET:
Posisi: {job.title}
Perusahaan: {job.company_name}
Lokasi: {job.location}
Deskripsi Lowongan:
{job.description}

DOKUMEN RESUME ASLI KANDIDAT:
{raw_resume_text}
"""

    # Daftar model dengan automatic fallback bila salah satu model terkena 429 quota exhaustion, 503, atau 404
    models_to_try = [
        "gemini-flash-lite-latest",
        "gemini-3.5-flash-lite",
        "gemini-3.1-flash-lite",
        "gemini-3.1-flash-lite-preview",
        "gemini-3.7-flash",
        "gemini-3.6-flash",
        "gemini-3.5-flash",
        "gemini-flash-latest",
    ]


    last_error = None
    for model_name in models_to_try:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=build_expert_ats_prompt(),
                    response_mime_type="application/json",
                    response_schema=OptimizedResumeResult,
                    temperature=0.15,
                ),
            )
            data = json.loads(response.text)
            return OptimizedResumeResult(**data)
        except Exception as e:
            last_error = e
            err_str = str(e)
            if any(code in err_str for code in ["429", "RESOURCE_EXHAUSTED", "404", "503", "UNAVAILABLE"]):
                print(f"[RESUME-BUILDER-A] Model {model_name} quota/error: {e}. Mencoba fallback berikutnya...")
                continue
            raise e

    raise last_error


async def build_resume_mode_b(req: ManualProfileBuildRequest, job: JobRequisition) -> OptimizedResumeResult:
    """Mode B: Pembuatan resume dari form profil terarah kandidat."""
    if not GEMINI_API_KEY:
        raise ValueError("GEMINI_API_KEY belum disetel di .env")

    client = genai.Client(api_key=GEMINI_API_KEY)

    prompt = f"""LOWONGAN PEKERJAAN TARGET:
Posisi: {job.title}
Perusahaan: {job.company_name}
Deskripsi Lowongan:
{job.description}

PROFIL & RIWAYAT KANDIDAT:
Nama: {req.contact.full_name}
Email: {req.contact.email}
Phone: {req.contact.phone}
Lokasi: {req.contact.location}
LinkedIn: {req.contact.linkedin_url}
GitHub: {req.contact.github_url}
Daftar Skills Input: {', '.join(req.skills_input)}

Cerita Pengalaman & Proyek Kandidat:
{req.background_text}
"""

    models_to_try = [
        "gemini-flash-lite-latest",
        "gemini-3.5-flash-lite",
        "gemini-3.1-flash-lite",
        "gemini-3.1-flash-lite-preview",
        "gemini-3.7-flash",
        "gemini-3.6-flash",
        "gemini-3.5-flash",
        "gemini-flash-latest",
    ]


    last_error = None
    for model_name in models_to_try:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=build_expert_ats_prompt(),
                    response_mime_type="application/json",
                    response_schema=OptimizedResumeResult,
                    temperature=0.15,
                ),
            )
            data = json.loads(response.text)
            return OptimizedResumeResult(**data)
        except Exception as e:
            last_error = e
            err_str = str(e)
            if any(code in err_str for code in ["429", "RESOURCE_EXHAUSTED", "404", "503", "UNAVAILABLE"]):
                print(f"[RESUME-BUILDER-B] Model {model_name} quota/error: {e}. Mencoba fallback berikutnya...")
                continue
            raise e

    raise last_error


async def generate_cover_letter_only(resume_data: FullResumeData, job: JobRequisition) -> str:
    """Regenerasi Cover Letter secara independen tanpa menyentuh resume."""
    if not GEMINI_API_KEY:
        raise ValueError("GEMINI_API_KEY belum disetel di .env")

    client = genai.Client(api_key=GEMINI_API_KEY)

    prompt = f"""Tuliskan Cover Letter profesional yang disesuaikan secara spesifik:

TARGET PERUSAHAAN: {job.company_name}
TARGET POSISI: {job.title}
DESKRIPSI LOWONGAN:
{job.description[:2000]}

PROFIL KANDIDAT:
Nama: {resume_data.contact.full_name}
Email: {resume_data.contact.email}
Lokasi: {resume_data.contact.location}
Keahlian: {', '.join(resume_data.technical_skills)}
Pengalaman Utama:
{json.dumps([{'role': exp.position, 'company': exp.company, 'bullets': exp.bullets} for exp in resume_data.work_experiences[:2]])}

Instruksi: Tulis Cover Letter 3-4 paragraf yang persuasif, personal, tidak klise, dan langsung relevan dengan kebutuhan peran tersebut. Kembalikan HANYA teks surat lengkap tanpa format JSON.
"""

    models_to_try = [
        "gemini-flash-lite-latest",
        "gemini-3.5-flash-lite",
        "gemini-3.1-flash-lite",
        "gemini-3.1-flash-lite-preview",
        "gemini-3.7-flash",
        "gemini-3.6-flash",
        "gemini-3.5-flash",
        "gemini-flash-latest",
    ]


    last_error = None
    for model_name in models_to_try:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=0.2,
                ),
            )
            return response.text.strip()
        except Exception as e:
            last_error = e
            err_str = str(e)
            if any(code in err_str for code in ["429", "RESOURCE_EXHAUSTED", "404", "503", "UNAVAILABLE"]):
                print(f"[RESUME-COVER-LETTER] Model {model_name} quota/error: {e}. Mencoba fallback berikutnya...")
                continue
            raise e

    raise last_error


def export_resume_to_docx_bytes(resume: FullResumeData, template: str = "HARVARD") -> io.BytesIO:
    """Membuat file DOCX resume ATS-friendly dengan pilihan template (HARVARD, MODERN, TECHNICAL).
    Menjamin reading order linear tanpa floating text-boxes atau tabel nested yang merusak parser ATS.
    """
    doc = Document()
    tpl = template.upper()

    safe_name = resume.contact.full_name or "Candidate"
    target_role = resume.target_role_title or "Professional Resume"
    
    # Metadata dokumen DOCX
    core_props = doc.core_properties
    core_props.title = f"Resume - {safe_name} ({target_role})"
    core_props.author = safe_name
    core_props.subject = f"Curriculum Vitae / ATS-Optimized Resume for {target_role}"
    core_props.category = "Curriculum Vitae"
    core_props.comments = "Created by Cakwe Career Radar & ATS Studio (ATS Compliant)"


    font_family = "Calibri"
    if tpl == "HARVARD":
        font_family = "Times New Roman"
    elif tpl == "TECHNICAL":
        font_family = "Arial"

    primary_color = RGBColor(15, 23, 42)
    accent_color = RGBColor(79, 70, 229) if tpl == "MODERN" else RGBColor(15, 23, 42)

    # Margin proporsional 0.6 inci
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # 1. Header Kontak
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.LEFT if tpl == "MODERN" else WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(resume.contact.full_name.upper())
    run_name.bold = True
    run_name.font.size = Pt(16)
    run_name.font.name = font_family
    run_name.font.color.rgb = accent_color

    contact_parts = []
    if resume.contact.email:
        contact_parts.append(resume.contact.email)
    if resume.contact.phone:
        contact_parts.append(resume.contact.phone)
    if resume.contact.location:
        contact_parts.append(resume.contact.location)
    if resume.contact.linkedin_url:
        contact_parts.append(resume.contact.linkedin_url)
    if resume.contact.github_url:
        contact_parts.append(resume.contact.github_url)

    p_contact = doc.add_paragraph(" | ".join(contact_parts))
    p_contact.alignment = WD_ALIGN_PARAGRAPH.LEFT if tpl == "MODERN" else WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(10)
    for r in p_contact.runs:
        r.font.size = Pt(9.5)
        r.font.name = font_family

    def add_section_heading(title: str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = font_family
        run.font.color.rgb = accent_color

    # 2. Professional Summary
    if resume.professional_summary:
        add_section_heading("Professional Summary")
        p_sum = doc.add_paragraph(resume.professional_summary)
        p_sum.paragraph_format.space_after = Pt(6)
        p_sum.paragraph_format.line_spacing = 1.15
        for r in p_sum.runs:
            r.font.size = Pt(10)
            r.font.name = font_family

    # 3. Technical Skills
    if resume.technical_skills:
        add_section_heading("Technical Skills")
        p_skills = doc.add_paragraph()
        p_skills.paragraph_format.space_after = Pt(6)
        r_skill_label = p_skills.add_run("Core Competencies: ")
        r_skill_label.bold = True
        r_skill_label.font.size = Pt(10)
        r_skill_label.font.name = font_family
        r_skills = p_skills.add_run(", ".join(resume.technical_skills))
        r_skills.font.size = Pt(10)
        r_skills.font.name = font_family

    # 4. Work Experience
    if resume.work_experiences:
        add_section_heading("Professional Experience")
        for exp in resume.work_experiences:
            p_exp_head = doc.add_paragraph()
            p_exp_head.paragraph_format.space_before = Pt(4)
            p_exp_head.paragraph_format.space_after = Pt(1)
            p_exp_head.paragraph_format.keep_with_next = True
            
            r_pos = p_exp_head.add_run(exp.position)
            r_pos.bold = True
            r_pos.font.size = Pt(10.5)
            r_pos.font.name = font_family

            r_comp = p_exp_head.add_run(f" | {exp.company}")
            r_comp.font.size = Pt(10.5)
            r_comp.font.name = font_family

            dates_str = f"{exp.start_date} – {exp.end_date}"
            if exp.location:
                dates_str += f" | {exp.location}"
            p_dates = doc.add_paragraph(dates_str)
            p_dates.paragraph_format.space_after = Pt(3)
            p_dates.paragraph_format.keep_with_next = True
            for r in p_dates.runs:
                r.font.size = Pt(9.5)
                r.italic = True
                r.font.name = font_family
                r.font.color.rgb = RGBColor(71, 85, 105)

            for b in exp.bullets:
                p_b = doc.add_paragraph(b, style="List Bullet")
                p_b.paragraph_format.space_before = Pt(0)
                p_b.paragraph_format.space_after = Pt(2)
                p_b.paragraph_format.line_spacing = 1.15
                for r in p_b.runs:
                    r.font.size = Pt(10)
                    r.font.name = font_family

    # 5. Education
    if resume.educations:
        add_section_heading("Education")
        for edu in resume.educations:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = Pt(2)
            p_edu.paragraph_format.space_after = Pt(1)
            r_deg = p_edu.add_run(f"{edu.degree} in {edu.field_of_study}")
            r_deg.bold = True
            r_deg.font.size = Pt(10)
            r_deg.font.name = font_family

            p_inst = doc.add_paragraph(f"{edu.institution} | {edu.graduation_year}")
            p_inst.paragraph_format.space_after = Pt(4)
            for r in p_inst.runs:
                r.font.size = Pt(9.5)
                r.font.name = font_family
                r.font.color.rgb = RGBColor(71, 85, 105)

    # 6. Certifications
    if resume.certifications:
        add_section_heading("Certifications")
        for cert in resume.certifications:
            p_cert = doc.add_paragraph(f"• {cert.name} — {cert.issuer} ({cert.year})")
            p_cert.paragraph_format.space_after = Pt(2)
            for r in p_cert.runs:
                r.font.size = Pt(9.5)
                r.font.name = font_family

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


def validate_resume_data(resume: FullResumeData) -> dict:
    """Audit lokal struktur resume untuk memastikan kelayakan ATS:
    - Keterbacaan dan kelengkapan informasi kontak
    - Keberadaan seksi esensial (Summary, Skills, Work Experience, Education)
    - Rasio bullet points terukur (formula dampak)
    - Peringatan jika terdapat placeholder teks yang belum diganti (seperti [X%])
    """
    issues = []
    strengths = []
    score = 100

    # 1. Cek Kontak
    contact = resume.contact
    contact_complete = bool(contact.full_name and contact.email and contact.location)
    if not contact_complete:
        issues.append("Informasi kontak belum lengkap (Nama, Email, dan Lokasi wajib diisi).")
        score -= 20
    else:
        strengths.append("Informasi kontak lengkap dan jelas.")

    # 2. Cek Summary
    has_summary = bool(resume.professional_summary and len(resume.professional_summary.split()) >= 15)
    if not has_summary:
        issues.append("Professional summary terlalu singkat atau kosong.")
        score -= 15
    else:
        strengths.append("Professional summary padat dan informatif.")

    # 3. Cek Skills
    has_skills = bool(resume.technical_skills and len(resume.technical_skills) >= 4)
    if not has_skills:
        issues.append("Daftar keahlian teknis minimal harus memiliki 4 item.")
        score -= 15
    else:
        strengths.append(f"{len(resume.technical_skills)} kata kunci keahlian teknis terstruktur rapi.")

    # 4. Cek Work Experiences
    has_experience = bool(resume.work_experiences and len(resume.work_experiences) > 0)
    quantified_count = 0
    total_bullets = 0
    placeholder_warnings = []

    if not has_experience:
        issues.append("Belum ada riwayat pengalaman kerja profesional.")
        score -= 30
    else:
        for exp in resume.work_experiences:
            for b in exp.bullets:
                total_bullets += 1
                # Cek jika ada angka atau persen
                if any(char.isdigit() for char in b):
                    quantified_count += 1
                if "[x" in b.lower() or "[n" in b.lower():
                    placeholder_warnings.append(f"Terdapat placeholder '{b[:40]}...' yang belum diisi angka riil.")

        quantified_ratio = (quantified_count / total_bullets) if total_bullets > 0 else 0
        if quantified_ratio < 0.4:
            issues.append("Tingkatkan bullet point dengan angka/metrik terukur (minimal 40% dari total bullet).")
            score -= 10
        else:
            strengths.append(f"{int(quantified_ratio * 100)}% bullet point menggunakan formula dampak terukur.")

        if placeholder_warnings:
            issues.extend(placeholder_warnings[:2])
            score -= 10

    # 5. Cek Education
    has_education = bool(resume.educations and len(resume.educations) > 0)
    if not has_education:
        issues.append("Informasi pendidikan atau sertifikasi profesional belum dicantumkan.")
        score -= 10

    total_words = len(resume.professional_summary.split()) + sum(len(b.split()) for exp in resume.work_experiences for b in exp.bullets)

    final_score = max(score, 20)

    return {
        "score": final_score,
        "is_valid": final_score >= 70,
        "issues": issues,
        "strengths": strengths,
        "word_count": total_words,
        "section_checks": {
            "contact_complete": contact_complete,
            "has_summary": has_summary,
            "has_skills": has_skills,
            "has_experience": has_experience,
            "has_education": has_education,
            "quantified_bullets_ratio": round((quantified_count / total_bullets) if total_bullets > 0 else 0, 2),
        }
    }


def export_cover_letter_to_docx_bytes(candidate_name: str, company_name: str, job_title: str, cover_letter_text: str) -> io.BytesIO:
    """Membuat file DOCX cover letter profesional."""
    doc = Document()
    core_props = doc.core_properties
    core_props.title = f"Cover Letter - {candidate_name} ({job_title} at {company_name})"
    core_props.author = candidate_name
    core_props.subject = f"Cover Letter Application for {job_title} at {company_name}"
    core_props.category = "Cover Letter"
    core_props.comments = "Created by Cakwe Career Radar & ATS Studio"

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    p_name = doc.add_paragraph()
    r_name = p_name.add_run(candidate_name.upper())
    r_name.bold = True
    r_name.font.size = Pt(14)
    r_name.font.name = "Calibri"

    p_target = doc.add_paragraph(f"Application for: {job_title} at {company_name}")
    p_target.paragraph_format.space_after = Pt(18)
    for r in p_target.runs:
        r.italic = True
        r.font.size = Pt(10)
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(71, 85, 105)

    # Paragraf-paragraf cover letter
    paragraphs = cover_letter_text.split("\n\n")
    for para in paragraphs:
        if para.strip():
            p = doc.add_paragraph(para.strip())
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.2
            for r in p.runs:
                r.font.size = Pt(11)
                r.font.name = "Calibri"

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


def export_resume_to_pdf_bytes(resume: FullResumeData, template: str = "HARVARD") -> io.BytesIO:
    """Membuat file PDF resume ATS-friendly dengan standar resmi Harvard/MIT Career Services."""
    import html
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buffer = io.BytesIO()
    # Margin standar Harvard Resume: 0.5 inci (36 pt)
    margin = 0.5 * inch
    page_width, page_height = letter
    content_width = page_width - (2 * margin)
    content_height = page_height - (2 * margin)

    safe_name = resume.contact.full_name or "Candidate"
    target_role = resume.target_role_title or "Professional Resume"

    doc = BaseDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=f"Resume - {safe_name} ({target_role})",
        author=safe_name,
        subject=f"Curriculum Vitae / ATS-Optimized Resume for {target_role}",
        creator="Cakwe Career Radar & ATS Studio",
        producer="ReportLab PDF Library (Harvard ATS Standard)",
    )
    # Gunakan zero-padding Frame agar posisi X dari Paragraph, HR line, dan Table presisi 100% rata kiri
    frame = Frame(
        margin,
        margin,
        content_width,
        content_height,
        id='normal',
        leftPadding=0,
        rightPadding=0,
        topPadding=0,
        bottomPadding=0,
    )
    doc.addPageTemplates([PageTemplate(id='First', frames=frame, pagesize=letter)])

    tpl = template.upper()
    if tpl == "MODERN":
        font_body = "Helvetica"
        font_bold = "Helvetica-Bold"
        font_italic = "Helvetica-Oblique"
        primary_color = colors.HexColor("#0f172a")
        accent_color = colors.HexColor("#312e81")
    elif tpl == "TECHNICAL":
        font_body = "Helvetica"
        font_bold = "Helvetica-Bold"
        font_italic = "Helvetica-Oblique"
        primary_color = colors.HexColor("#090d16")
        accent_color = colors.HexColor("#0f172a")
    else:  # HARVARD (Serif standar resmi)
        font_body = "Times-Roman"
        font_bold = "Times-Bold"
        font_italic = "Times-Italic"
        primary_color = colors.HexColor("#111827")
        accent_color = colors.HexColor("#111827")

    styles = getSampleStyleSheet()

    # Tipografi standar resmi Harvard / Wall Street resume
    name_style = ParagraphStyle(
        'HarvardName',
        parent=styles['Normal'],
        fontName=font_bold,
        fontSize=18,
        leading=21,
        alignment=1,  # Center
        textColor=primary_color,
    )

    contact_style = ParagraphStyle(
        'HarvardContact',
        parent=styles['Normal'],
        fontName=font_body,
        fontSize=9,
        leading=12,
        alignment=1,  # Center
        textColor=colors.HexColor("#374151"),
    )

    section_heading_style = ParagraphStyle(
        'HarvardSectionHeading',
        parent=styles['Normal'],
        fontName=font_bold,
        fontSize=10.5,
        leading=13,
        textColor=accent_color,
        spaceBefore=0,
        spaceAfter=1,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        'HarvardBody',
        parent=styles['Normal'],
        fontName=font_body,
        fontSize=9.5,
        leading=13,
        textColor=primary_color,
    )

    table_left_style = ParagraphStyle(
        'TableLeft',
        parent=styles['Normal'],
        fontName=font_body,
        fontSize=9.5,
        leading=13,
        leftIndent=0,
        firstLineIndent=0,
        textColor=primary_color,
    )

    table_right_style = ParagraphStyle(
        'TableRight',
        parent=styles['Normal'],
        fontName=font_italic,
        fontSize=9,
        leading=13,
        alignment=2,  # Rata Kanan
        textColor=colors.HexColor("#4b5563"),
    )

    bullet_style = ParagraphStyle(
        'HarvardBullet',
        parent=styles['Normal'],
        fontName=font_body,
        fontSize=9.2,
        leading=12.2,
        leftIndent=11,
        firstLineIndent=-11,
        spaceAfter=1.8,
        textColor=primary_color,
    )

    story = []

    # 1. Header Nama & Kontak (Centered Harvard Standard)
    name_escaped = html.escape(resume.contact.full_name.upper())
    story.append(Paragraph(name_escaped, name_style))
    story.append(Spacer(1, 2))

    contact_items = []
    if resume.contact.email:
        contact_items.append(html.escape(resume.contact.email))
    if resume.contact.phone:
        contact_items.append(html.escape(resume.contact.phone))
    if resume.contact.location:
        contact_items.append(html.escape(resume.contact.location))
    if resume.contact.linkedin_url:
        contact_items.append(html.escape(resume.contact.linkedin_url))
    if resume.contact.github_url:
        contact_items.append(html.escape(resume.contact.github_url))

    story.append(Paragraph(" | ".join(contact_items), contact_style))
    story.append(Spacer(1, 4))

    def add_section(title: str):
        story.append(Paragraph(title.upper(), section_heading_style))
        story.append(HRFlowable(
            width="100%",
            thickness=0.75,
            color=colors.HexColor("#1f2937"),
            spaceBefore=1,
            spaceAfter=3,
            hAlign='LEFT'
        ))

    # 2. Professional Summary
    if resume.professional_summary:
        add_section("Professional Summary")
        story.append(Paragraph(html.escape(resume.professional_summary), body_style))
        story.append(Spacer(1, 4))

    # 3. Technical Skills
    if resume.technical_skills:
        add_section("Technical Skills")
        skills_esc = ", ".join(html.escape(str(s)) for s in resume.technical_skills)
        story.append(Paragraph(f"<b>Core Competencies:</b> {skills_esc}", body_style))
        story.append(Spacer(1, 4))

    # 4. Professional Experience
    if resume.work_experiences:
        add_section("Professional Experience")
        for idx, exp in enumerate(resume.work_experiences):
            pos_esc = html.escape(exp.position)
            comp_esc = html.escape(exp.company)
            dates_esc = html.escape(f"{exp.start_date} – {exp.end_date}")
            loc_esc = f" | {html.escape(exp.location)}" if exp.location else ""

            # Table tanpa margin/padding sisi luar (0 padding) agar persis rata kiri garis HR
            left_p = Paragraph(f"<b>{pos_esc}</b>, <i>{comp_esc}</i>", table_left_style)
            right_p = Paragraph(f"{dates_esc}{loc_esc}", table_right_style)

            tbl = Table([[left_p, right_p]], colWidths=[content_width * 0.65, content_width * 0.35])
            tbl.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 1),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ]))
            story.append(tbl)

            for b in exp.bullets:
                clean_b = html.escape(b.lstrip("•-* ").strip())
                story.append(Paragraph(f"•&nbsp;&nbsp;{clean_b}", bullet_style))
            story.append(Spacer(1, 2.5))

    # 5. Education
    if resume.educations:
        add_section("Education")
        for edu in resume.educations:
            inst_esc = html.escape(edu.institution)
            deg_esc = html.escape(edu.degree)
            field_esc = html.escape(edu.field_of_study)
            year_esc = html.escape(str(edu.graduation_year))

            left_edu = Paragraph(f"<b>{inst_esc}</b> — {deg_esc} in {field_esc}", table_left_style)
            right_edu = Paragraph(year_esc, table_right_style)

            tbl_edu = Table([[left_edu, right_edu]], colWidths=[content_width * 0.8, content_width * 0.2])
            tbl_edu.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 1),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
            ]))
            story.append(tbl_edu)
        story.append(Spacer(1, 3))

    # 6. Certifications
    if resume.certifications:
        add_section("Certifications")
        for cert in resume.certifications:
            if hasattr(cert, "name"):
                name_esc = html.escape(cert.name)
                issuer_esc = f" — {html.escape(str(cert.issuer))}" if getattr(cert, "issuer", None) else ""
                year_esc = html.escape(str(cert.year)) if getattr(cert, "year", None) else ""

                # Gunakan bullet konsisten persis di kiri teks sertifikasi
                left_cert = Paragraph(f"•&nbsp;&nbsp;<b>{name_esc}</b>{issuer_esc}", table_left_style)
                right_cert = Paragraph(year_esc, table_right_style)

                tbl_cert = Table([[left_cert, right_cert]], colWidths=[content_width * 0.82, content_width * 0.18])
                tbl_cert.setStyle(TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 0),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                    ('TOPPADDING', (0, 0), (-1, -1), 1),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
                ]))
                story.append(tbl_cert)
            else:
                story.append(Paragraph(f"•&nbsp;&nbsp;{html.escape(str(cert))}", bullet_style))

    doc.build(story)
    buffer.seek(0)
    return buffer


def export_cover_letter_to_pdf_bytes(candidate_name: str, company_name: str, job_title: str, cover_letter_text: str) -> io.BytesIO:
    """Membuat file PDF cover letter profesional 1 halaman (reportlab)."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title=f"Cover Letter - {candidate_name} ({job_title} at {company_name})",
        author=candidate_name,
        subject=f"Cover Letter Application for {job_title} at {company_name}",
        creator="Cakwe Career Radar & ATS Studio",
        producer="ReportLab PDF Library",
    )

    styles = getSampleStyleSheet()

    name_style = ParagraphStyle(
        'CoverName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0f172a"),
    )

    meta_style = ParagraphStyle(
        'CoverMeta',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#475569"),
    )

    body_style = ParagraphStyle(
        'CoverBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=10,
    )

    story = []
    story.append(Paragraph(candidate_name.upper(), name_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph(f"Application for: <b>{job_title}</b> at <b>{company_name}</b>", meta_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e2e8f0"), spaceBefore=2, spaceAfter=14))

    for p in cover_letter_text.split("\n\n"):
        p_clean = p.strip()
        if p_clean:
            story.append(Paragraph(p_clean, body_style))

    story.append(Spacer(1, 15))
    story.append(Paragraph("Hormat saya,", body_style))
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"<b>{candidate_name}</b>", body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer
